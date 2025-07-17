import os
import sys

import pyodbc
from openpyxl import Workbook


def main():
    if len(sys.argv) < 2:
        print('Usage: python export_instruments.py <path_to_mdb>')
        return

    mdb_path = sys.argv[1]
    if not os.path.exists(mdb_path):
        print(f'MDB file not found: {mdb_path}')
        return

    conn_str = (
        r'DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};'
        f'DBQ={mdb_path};'
    )
    try:
        conn = pyodbc.connect(conn_str, autocommit=True)
        cursor = conn.cursor()
        table_names = [row.table_name for row in cursor.tables(tableType='TABLE')]
        if 'Instruments' not in table_names:
            print('Instruments table not found in the MDB file.')
            return
        query = (
            "SELECT ID, Tag, FullDescription, EGULow, EGUHigh, RawLow, RawHigh, "
            "HALM_EN, HALM_SP, HALM_DB, HALM_DLY, "
            "HWARN_EN, HWARN_SP, HWARN_DB, HWARN_DLY, "
            "LALM_EN, LALM_SP, LALM_DB, LALM_DLY, "
            "LWARN_EN, LWARN_SP, LWARN_DB, LWARN_DLY, "
            "DigitalInput, DigitalOutput, AnalogInput, AnalogOutput "
            "FROM Instruments WHERE Type='IO' AND Tag <> '' AND Tag IS NOT NULL"
        )
        cursor.execute(query)
        rows = cursor.fetchall()
    except pyodbc.Error as e:
        print(f'Error accessing MDB file: {e}')
        return
    finally:
        try:
            conn.close()
        except Exception:
            pass

    output_path = input('Enter path to save Excel file: ').strip()
    if not output_path:
        print('No output path provided')
        return

    header = [
        'ID',
        'Tag',
        'FullDescription',
        'EGULow',
        'EGUHigh',
        'RawLow',
        'RawHigh',
        'HALM_EN',
        'HALM_SP',
        'HALM_DB',
        'HALM_DLY',
        'HWARN_EN',
        'HWARN_SP',
        'HWARN_DB',
        'HWARN_DLY',
        'LALM_EN',
        'LALM_SP',
        'LALM_DB',
        'LALM_DLY',
        'LWARN_EN',
        'LWARN_SP',
        'LWARN_DB',
        'LWARN_DLY',
    ]

    categories = {
        'DigitalInput': [],
        'DigitalOutput': [],
        'AnalogInput': [],
        'AnalogOutput': [],
    }

    for row in rows:
        if getattr(row, 'DigitalInput', False):
            categories['DigitalInput'].append(row)
        elif getattr(row, 'DigitalOutput', False):
            categories['DigitalOutput'].append(row)
        elif getattr(row, 'AnalogInput', False):
            categories['AnalogInput'].append(row)
        elif getattr(row, 'AnalogOutput', False):
            categories['AnalogOutput'].append(row)

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    for name, rows_list in categories.items():
        ws = wb.create_sheet(name)
        ws.append(header)
        for row in rows_list:
            ws.append(row[:len(header)])

    wb.save(output_path)
    print(f'Exported {sum(len(v) for v in categories.values())} rows to {output_path}')


if __name__ == '__main__':
    main()
